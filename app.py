import re
import json
import html
from urllib.request import Request, urlopen
from urllib.parse import urlencode, quote
from urllib.error import HTTPError, URLError
from io import BytesIO
from xml.etree import ElementTree as ET
from email.utils import parsedate_to_datetime
from urllib.parse import urlparse
from datetime import date, datetime, timedelta, timezone
from collections import Counter
import streamlit as st
import fitz
from docx import Document
import pandas as pd

st.set_page_config(page_title='AI Job Finder', page_icon='💼', layout='wide')

SKILLS = ['Leadership','Strategic Planning','Marketing','Data Analysis','Logistics','Compliance','Project Management','Operations Management','Budget Management','P&L Management','Advertising Sales','Digital Strategy','Audience Development','Community Relations','Business Development','Team Management','Revenue Growth','Financial Planning','Change Management','Circulation','Editorial Leadership','Sales Management','Vendor Management']
TITLES = ['Group Newspaper Publisher','Group Publisher','Newspaper Publisher','Publisher','Operations Director','Director of Operations','General Manager','Regional Manager','Operations Manager','Project Manager','Sales Director','Vice President']
STATUSES = ['Saved','Applied','Interviewing','Offer','Rejected','Skipped','Withdrawn']
PUBLISHER_AREAS = ['Multi-publication or multi-site leadership','Revenue, profitability, and P&L oversight','Advertising and sales leadership','Digital publishing and subscriptions','Editorial and cross-functional leadership','Budgeting, forecasting, and financial planning','Audience growth and marketing','Community and advertiser partnerships','Circulation, distribution, and logistics','Talent development and organizational change']

# These eight responsibilities were explicitly confirmed by the user.
# They remain editable; session state is not permanent storage.
VERIFICATION_AREAS = {
    'Full profit-and-loss (P&L) responsibility': 'Financial & P&L leadership',
    'Managing multiple business locations or publications': 'Multi-site business leadership',
    'Hiring, supervising, and developing department managers': 'Talent & organizational change',
    'Negotiating major contracts and vendor agreements': 'Negotiation & vendor management',
    'Developing annual budgets and financial forecasts': 'Financial & P&L leadership',
    'Leading revenue growth and business transformation': 'Revenue growth & transformation',
    'Overseeing distribution and logistics operations': 'Distribution & logistics',
    'Managing digital products and revenue': 'Digital transformation',
}
INITIAL_VERIFIED = list(VERIFICATION_AREAS)

for key, default in {'jobs': [], 'resume_text': '', 'confirmed_titles': [], 'analysis': None,
                     'approved_suggestions': [], 'verified_duties': INITIAL_VERIFIED.copy()}.items():
    if key not in st.session_state:
        st.session_state[key] = default


def extract_resume_text(file):
    data = file.getvalue()
    name = file.name.lower()
    if name.endswith('.pdf'):
        with fitz.open(stream=data, filetype='pdf') as doc:
            return '\n'.join(page.get_text() for page in doc)
    if name.endswith('.docx'):
        doc = Document(BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines.append(cell.text)
        return '\n'.join(lines)
    if name.endswith('.txt'):
        return data.decode('utf-8-sig', errors='replace')
    return ''


def keyword_matches(text, options):
    return [item for item in options if re.search(r'\b' + re.escape(item) + r'\b', text, re.I)]


# A local, transparent reference library. These are examples of common publisher
# responsibilities, not assertions about the user's actual experience.
COMPETENCIES = [
    ('Financial & P&L leadership', ['profitability', 'profit and loss', 'p&l', 'budget', 'forecast', 'cash flow', 'financial performance'], 'Financial leadership', 'Regional General Manager'),
    ('Multi-site business leadership', ['multiple newspapers', 'multiple publications', 'group publisher', 'group newspaper publisher', 'multi-site', 'multi-location', 'regional operations'], 'Multi-location leadership', 'Regional Operations Director'),
    ('Revenue growth & transformation', ['revenue growth', 'revenue increase', 'business transformation', 'turnaround', 'restructur', 'growth strategy'], 'Revenue and change leadership', 'General Manager'),
    ('Advertising & sales leadership', ['advertising', 'ad sales', 'sales team', 'sales director', 'sales revenue', 'account executives'], 'Commercial leadership', 'Director of Sales'),
    ('Digital transformation', ['digital revenue', 'digital marketing', 'online advertising', 'digital subscription', 'digital strategy', 'website'], 'Digital business growth', 'Digital Strategy Director'),
    ('Editorial & cross-functional leadership', ['editorial', 'managing editor', 'newsroom', 'editor', 'cross-functional'], 'Cross-functional management', 'General Manager'),
    ('Audience & marketing development', ['circulation', 'audience', 'subscriber', 'subscription', 'marketing', 'readership'], 'Customer growth', 'Marketing Director'),
    ('Distribution & logistics', ['distribution', 'delivery', 'route', 'logistics', 'circulation director'], 'Operations optimization', 'Distribution Director'),
    ('Talent & organizational change', ['recruit', 'hiring', 'training', 'cross-training', 'reorganization', 'staff development', 'employee'], 'People leadership', 'Director of Operations'),
    ('Community partnerships', ['community', 'chamber', 'partnership', 'advertiser relationships', 'stakeholder'], 'Stakeholder relations', 'Community Relations Director'),
    ('Negotiation & vendor management', ['negotiat', 'vendor', 'contract', 'supplier', 'insurance premium'], 'Negotiation', 'Procurement Director'),
]


def analyze_locally(resume):
    lines = [line.strip() for line in resume.splitlines() if line.strip()]
    documented, verify, roles, suggestions = [], [], {}, []
    for skill, phrases, transferable, role in COMPETENCIES:
        evidence = []
        for line in lines:
            if any(re.search(r'(?<!\w)' + re.escape(term) + r'(?!\w)', line, re.I) for term in phrases):
                evidence.append(line[:350])
            if len(evidence) == 2:
                break
        if evidence:
            documented.append({'skill': skill, 'evidence': evidence, 'transferable': transferable})
            roles.setdefault(role, []).append(transferable)
            suggestions.append({'skill': skill, 'source': evidence[0],
                'wording': f'Experience in {transferable.lower()}: {evidence[0]}',
                'note': 'Draft based on resume wording; edit for accuracy before using.'})
        else:
            verify.append(skill)
    return {'documented': documented, 'verify': verify, 'roles': roles, 'suggestions': suggestions}

# Career matches use the same publisher competency evidence as Career Analysis.
# Scores are transparent overlap indicators, not employment eligibility predictions.
CAREER_PATHS = [
    {'title': 'Regional Operations Director', 'industry': 'Multi-location businesses', 'core': ['Multi-site business leadership', 'Financial & P&L leadership', 'Talent & organizational change'], 'additional': ['Distribution & logistics', 'Negotiation & vendor management'], 'focus': 'Emphasize locations led, budgets, operational improvements and team size.'},
    {'title': 'General Manager', 'industry': 'Services, media, retail and local businesses', 'core': ['Financial & P&L leadership', 'Advertising & sales leadership', 'Talent & organizational change'], 'additional': ['Multi-site business leadership', 'Revenue growth & transformation'], 'focus': 'Highlight overall business ownership, revenue, profitability and leadership results.'},
    {'title': 'Director of Business Development', 'industry': 'Business services and B2B organizations', 'core': ['Advertising & sales leadership', 'Community partnerships', 'Financial & P&L leadership'], 'additional': ['Revenue growth & transformation', 'Negotiation & vendor management'], 'focus': 'Show account growth, partnerships, pricing strategy and measurable sales outcomes.'},
    {'title': 'Director of Marketing', 'industry': 'Marketing, media and consumer services', 'core': ['Audience & marketing development', 'Digital transformation', 'Advertising & sales leadership'], 'additional': ['Community partnerships', 'Financial & P&L leadership'], 'focus': 'Show audience growth, campaign results, digital revenue and customer acquisition.'},
    {'title': 'Distribution Operations Director', 'industry': 'Logistics and delivery services', 'core': ['Distribution & logistics', 'Talent & organizational change', 'Financial & P&L leadership'], 'additional': ['Multi-site business leadership', 'Negotiation & vendor management'], 'focus': 'Show delivery efficiencies, route changes, service levels and cost savings.'},
    {'title': 'Director of Community Partnerships', 'industry': 'Nonprofits, education and civic organizations', 'core': ['Community partnerships', 'Talent & organizational change', 'Audience & marketing development'], 'additional': ['Advertising & sales leadership', 'Financial & P&L leadership'], 'focus': 'Describe stakeholder relationships, public-facing leadership and partnership outcomes.'},
    {'title': 'Digital Business Director', 'industry': 'Digital media and subscription businesses', 'core': ['Digital transformation', 'Audience & marketing development', 'Financial & P&L leadership'], 'additional': ['Advertising & sales leadership', 'Editorial & cross-functional leadership'], 'focus': 'Show digital product initiatives, subscription or advertising growth, and cross-team execution.'},
    {'title': 'Chief Operating Officer (smaller organization)', 'industry': 'Small and midsize organizations', 'core': ['Financial & P&L leadership', 'Multi-site business leadership', 'Talent & organizational change'], 'additional': ['Negotiation & vendor management', 'Revenue growth & transformation'], 'focus': 'Document executive decision-making, business scale, direct reports and financial accountability.'},
]


def verified_competencies():
    return set(VERIFICATION_AREAS[duty] for duty in st.session_state.verified_duties)


def verification_editor():
    st.subheader('Experience Verification')
    st.write('These eight publisher responsibilities were preselected because you confirmed them. '
             'Uncheck anything you want to remove. Changes affect career matching immediately.')
    with st.form('experience_verification_form'):
        selected = []
        for duty in VERIFICATION_AREAS:
            if st.checkbox(duty, value=duty in st.session_state.verified_duties, key='verify_' + str(list(VERIFICATION_AREAS).index(duty))):
                selected.append(duty)
        if st.form_submit_button('Save verified experience'):
            st.session_state.verified_duties = selected
            st.session_state.pop('career_matches', None)
            st.success('Verified responsibilities updated for this session.')
    st.caption('Confirmed experience is based on your approval, not inferred from keywords. '
               'It is not automatically added to your resume.')


def get_career_matches(resume):
    analysis = analyze_locally(resume)
    evidence_by_skill = {item['skill']: item['evidence'] for item in analysis['documented']}
    verified = verified_competencies()
    results = []
    for path in CAREER_PATHS:
        all_skills = path['core'] + path['additional']
        matched = [skill for skill in all_skills if skill in evidence_by_skill or skill in verified]
        unconfirmed = [skill for skill in all_skills if skill not in matched]
        # Core requirements count twice as much as additional competencies.
        total_weight = 2 * len(path['core']) + len(path['additional'])
        matched_weight = sum(2 if skill in path['core'] else 1 for skill in matched)
        score = round(100 * matched_weight / total_weight)
        results.append({**path, 'score': score, 'matched': matched,
                        'unconfirmed': unconfirmed, 'evidence': evidence_by_skill,
                        'verified': [skill for skill in matched if skill in verified]})
    return sorted(results, key=lambda x: (-x['score'], x['title']))


# Free public remote-job feeds. Cache requests to avoid unnecessary provider traffic.
# Search APIs query the provider directly, rather than filtering only a small
# downloaded feed. Listings must be independently checked before applying.
SEARCH_TERMS = [
    'operations', 'business development', 'general manager', 'marketing',
    'partnerships', 'logistics', 'chief operating officer', 'publisher'
]


def fetch_json(url):
    request = Request(url, headers={
        'User-Agent': 'Mozilla/5.0 (compatible; PersonalCareerResearch/1.0)',
        'Accept': 'application/json'})
    with urlopen(request, timeout=16) as response:
        return json.load(response)


MAX_POSTING_AGE_DAYS = 45


def parse_posted_date(value):
    """Return a date for common API date formats; unknown dates stay unknown."""
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        try:
            timestamp = float(value)
            if timestamp > 1e11:  # Unix milliseconds
                timestamp /= 1000
            return datetime.fromtimestamp(timestamp, tz=timezone.utc).date()
        except (ValueError, OverflowError, OSError):
            return None
    raw = str(value).strip()
    try:
        return datetime.fromisoformat(raw.replace('Z', '+00:00')).date()
    except ValueError:
        pass
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        pass
    try:
        return parsedate_to_datetime(raw).date()
    except (ValueError, TypeError, OverflowError):
        return None


def is_recent_posting(job, today=None):
    """Only accept known publication dates from the last 45 days."""
    today = today or datetime.now(timezone.utc).date()
    published = parse_posted_date(job.get('date'))
    return published is not None and today - timedelta(days=MAX_POSTING_AGE_DAYS) <= published <= today


def standard_job(title, company, url, location, description, source,
                 remote=True, salary='', posted='', eligible=True):
    if not (title and company and str(url).startswith('https://')):
        return None
    return {'title': str(title).strip(), 'company': str(company).strip(),
            'url': str(url), 'location': str(location or 'Not specified'),
            'remote': bool(remote), 'salary': str(salary or ''),
            'date': parse_posted_date(posted).isoformat() if parse_posted_date(posted) else '',
            'description': html_to_text(str(description or ''))[:6000],
            'source': source, 'eligible': eligible}


US_STATES = {
    'AL':'Alabama','AK':'Alaska','AZ':'Arizona','AR':'Arkansas','CA':'California',
    'CO':'Colorado','CT':'Connecticut','DE':'Delaware','FL':'Florida','GA':'Georgia',
    'HI':'Hawaii','ID':'Idaho','IL':'Illinois','IN':'Indiana','IA':'Iowa','KS':'Kansas',
    'KY':'Kentucky','LA':'Louisiana','ME':'Maine','MD':'Maryland','MA':'Massachusetts',
    'MI':'Michigan','MN':'Minnesota','MS':'Mississippi','MO':'Missouri','MT':'Montana',
    'NE':'Nebraska','NV':'Nevada','NH':'New Hampshire','NJ':'New Jersey','NM':'New Mexico',
    'NY':'New York','NC':'North Carolina','ND':'North Dakota','OH':'Ohio','OK':'Oklahoma',
    'OR':'Oregon','PA':'Pennsylvania','RI':'Rhode Island','SC':'South Carolina',
    'SD':'South Dakota','TN':'Tennessee','TX':'Texas','UT':'Utah','VT':'Vermont',
    'VA':'Virginia','WA':'Washington','WV':'West Virginia','WI':'Wisconsin','WY':'Wyoming',
    'DC':'District of Columbia'
}


def states_in_location(location):
    """Identify explicit U.S. states only; do not guess from vague locations."""
    text = str(location or '')
    found = set()
    for code, name in US_STATES.items():
        if re.search(r'(?<![A-Za-z])' + re.escape(name) + r'(?![A-Za-z])', text, re.I):
            found.add(code)
        # Abbreviations only after a comma, or as a standalone location.
        if re.search(r',\s*' + code + r'(?:\b|$)', text) or text.strip() == code:
            found.add(code)
    return sorted(found)


def muse_location(locations):
    names = [x.get('name','') for x in (locations or []) if isinstance(x, dict)]
    return ' / '.join(n for n in names if n) or 'Not specified'


def muse_remote(location):
    return bool(re.search(r'\b(remote|flexible|work from home|anywhere|virtual)\b', location, re.I))


def muse_us_eligible(location):
    return bool(states_in_location(location) or re.search(
        r'\b(United States|USA|U\.S\.|US Remote|Remote - US)\b', location, re.I))


@st.cache_data(ttl=3600, show_spinner=False)
def load_muse_jobs():
    """The Muse offers public location-tagged jobs, including non-remote roles."""
    jobs, errors = [], []
    # Public API: management and senior-level pages; small bounded request count.
    for level in ('management', 'Senior Level'):
        for page in range(0, 8):
            try:
                url = 'https://www.themuse.com/api/public/jobs?' + urlencode(
                    {'page': page, 'level': level})
                payload = fetch_json(url)
                records = payload.get('results', [])
                if not records:
                    break
                for rec in records:
                    location = muse_location(rec.get('locations'))
                    is_remote = muse_remote(location)
                    eligible = muse_us_eligible(location)
                    company = rec.get('company') or {}
                    company_name = company.get('name', '') if isinstance(company, dict) else str(company)
                    refs = rec.get('refs') or {}
                    link = refs.get('landing_page') or refs.get('external_apply_url') or ''
                    job = standard_job(rec.get('name'), company_name, link,
                        location, rec.get('contents', ''), 'The Muse', is_remote,
                        '', rec.get('publication_date'), eligible)
                    if job:
                        jobs.append(job)
            except (HTTPError, URLError, TimeoutError, ValueError, OSError, TypeError) as exc:
                errors.append(f'The Muse {level} page {page}: {type(exc).__name__}')
                break
    return jobs, errors


@st.cache_data(ttl=3600, show_spinner=False)
def load_public_jobs():
    jobs, errors, counts = [], [], {}
    muse_jobs, muse_errors = load_muse_jobs()
    jobs.extend(muse_jobs)
    errors.extend(muse_errors)
    counts['The Muse'] = len(muse_jobs)
    # Jobicy publishes recent remote jobs. Its USA filter is supplied by Jobicy.
    try:
        payload = fetch_json('https://jobicy.com/api/v2/remote-jobs?count=200&geo=usa')
        for rec in payload.get('jobs', []):
            job = standard_job(rec.get('jobTitle'), rec.get('companyName'),
                rec.get('url'), rec.get('jobGeo') or 'USA / Remote',
                rec.get('jobDescription'), 'Jobicy', True,
                rec.get('annualSalaryMin') or '', rec.get('pubDate'), True)
            if job: jobs.append(job)
        counts['Jobicy'] = len([j for j in jobs if j['source'] == 'Jobicy'])
    except (HTTPError, URLError, TimeoutError, ValueError, OSError, TypeError) as exc:
        errors.append(f'Jobicy could not be loaded ({type(exc).__name__}).')

    # Himalayas searches remotely hosted jobs by keyword and country. Searching
    # provider-side prevents executive titles from being lost in feed sampling.
    for term in SEARCH_TERMS:
        try:
            url = 'https://himalayas.app/jobs/api/search?' + urlencode(
                {'q': term, 'country': 'US', 'sort': 'recent', 'page': 1})
            payload = fetch_json(url)
            for rec in payload.get('jobs', []):
                restrictions = rec.get('locationRestrictions') or []
                location = ', '.join(str(x) for x in restrictions) if restrictions else 'Worldwide / Remote'
                low, high, currency = rec.get('minSalary'), rec.get('maxSalary'), rec.get('currency') or ''
                salary = f'{currency} {low:,}–{high:,} / {rec.get("salaryPeriod") or "year"}' if isinstance(low, (int,float)) and isinstance(high,(int,float)) else ''
                job = standard_job(rec.get('title'), rec.get('companyName'),
                    rec.get('applicationLink'), location,
                    rec.get('description') or rec.get('excerpt'), 'Himalayas',
                    True, salary, rec.get('pubDate'), True)
                if job: jobs.append(job)
        except (HTTPError, URLError, TimeoutError, ValueError, OSError, TypeError) as exc:
            errors.append(f'Himalayas search for {term} failed ({type(exc).__name__}).')
    counts['Himalayas'] = len([j for j in jobs if j['source'] == 'Himalayas'])

    # Remotive supplements those sources. Location restrictions are interpreted
    # conservatively, rather than assuming every remote job accepts U.S. applicants.
    try:
        payload = fetch_json('https://remotive.com/api/remote-jobs')
        for rec in payload.get('jobs', []):
            location = str(rec.get('candidate_required_location') or 'Unspecified')
            eligible = any(x in location.lower() for x in
                ['usa', 'united states', 'u.s.', 'north america', 'worldwide', 'anywhere', 'global'])
            job = standard_job(rec.get('title'), rec.get('company_name'),
                rec.get('url'), location, rec.get('description'), 'Remotive',
                True, rec.get('salary'), rec.get('publication_date'), eligible)
            if job: jobs.append(job)
        counts['Remotive'] = len([j for j in jobs if j['source'] == 'Remotive'])
    except (HTTPError, URLError, TimeoutError, ValueError, OSError, TypeError) as exc:
        errors.append(f'Remotive could not be loaded ({type(exc).__name__}).')
    unique = {}
    for job in jobs:
        unique.setdefault(job['url'].split('?')[0], job)
    return list(unique.values()), errors, counts

# Additional public employer career feeds. These APIs are employer-specific;
# tokens are entered by the user rather than guessed or scraped.
@st.cache_data(ttl=3600, show_spinner=False)
def load_employer_boards(greenhouse_tokens, lever_tokens):
    jobs, errors, counts = [], [], {}
    for token in greenhouse_tokens[:12]:
        try:
            payload = fetch_json(f'https://boards-api.greenhouse.io/v1/boards/{quote(token)}/jobs?content=true')
            records = payload.get('jobs', [])
            for rec in records:
                location = (rec.get('location') or {}).get('name', 'Not specified')
                remote = muse_remote(location) or muse_remote(rec.get('title', ''))
                job = standard_job(rec.get('title'), token, rec.get('absolute_url'), location,
                    rec.get('content'), 'Greenhouse: ' + token, remote, '',
                    rec.get('first_published') or rec.get('updated_at'),
                    muse_us_eligible(location))
                # updated_at is a modification date, NOT a publication date.
                # Never treat it as proof of a recent posting.
                if job and rec.get('first_published'):
                    jobs.append(job)
            counts['Greenhouse: ' + token] = sum(j['source'] == 'Greenhouse: ' + token for j in jobs)
        except (HTTPError, URLError, TimeoutError, ValueError, OSError, TypeError) as exc:
            errors.append(f'Greenhouse board {token}: {type(exc).__name__}')
    for token in lever_tokens[:12]:
        try:
            records = fetch_json(f'https://api.lever.co/v0/postings/{quote(token)}?mode=json&limit=100')
            for rec in records:
                categories = rec.get('categories') or {}
                location = categories.get('location') or 'Not specified'
                workplace = categories.get('workplaceType') or ''
                remote = workplace.lower() == 'remote' or muse_remote(location)
                job = standard_job(rec.get('text'), token, rec.get('hostedUrl'), location,
                    (rec.get('descriptionPlain') or '') + ' ' + (rec.get('additionalPlain') or ''),
                    'Lever: ' + token, remote, '', rec.get('createdAt'),
                    muse_us_eligible(location))
                if job: jobs.append(job)
            counts['Lever: ' + token] = sum(j['source'] == 'Lever: ' + token for j in jobs)
        except (HTTPError, URLError, TimeoutError, ValueError, OSError, TypeError) as exc:
            errors.append(f'Lever board {token}: {type(exc).__name__}')
    return jobs, errors, counts


@st.cache_data(ttl=3600, show_spinner=False)
def load_usajobs(api_key, email):
    if not api_key or not email:
        return [], [], {}
    jobs, errors = [], []
    for keyword in ['program manager', 'operations manager', 'business development',
                    'public affairs director', 'administrative officer']:
        try:
            url = 'https://data.usajobs.gov/api/search?' + urlencode(
                {'Keyword': keyword, 'ResultsPerPage': 100})
            request = Request(url, headers={'Host': 'data.usajobs.gov',
                'User-Agent': email, 'Authorization-Key': api_key,
                'Accept': 'application/json'})
            with urlopen(request, timeout=20) as response:
                payload = json.load(response)
            for item in (payload.get('SearchResult') or {}).get('SearchResultItems', []):
                rec = item.get('MatchedObjectDescriptor') or {}
                detail = rec.get('UserArea', {}).get('Details', {})
                locs = rec.get('PositionLocation') or []
                location = ' / '.join(x.get('LocationName', '') for x in locs) or 'United States'
                description = ' '.join(str(detail.get(k) or '') for k in
                    ['JobSummary', 'MajorDuties', 'QualificationSummary'])
                job = standard_job(rec.get('PositionTitle'), rec.get('OrganizationName'),
                    rec.get('PositionURI'), location, description, 'USAJOBS',
                    bool(re.search(r'\bremote\b', location, re.I)), '',
                    rec.get('PublicationStartDate'), True)
                if job: jobs.append(job)
        except (HTTPError, URLError, TimeoutError, ValueError, OSError, TypeError) as exc:
            errors.append(f'USAJOBS ({keyword}): {type(exc).__name__}')
    return jobs, errors, {'USAJOBS': len(jobs)}


def parse_board_tokens(value):
    tokens = [x.strip().lower() for x in re.split(r'[,\n;]+', value or '')]
    return tuple(dict.fromkeys(x for x in tokens if re.fullmatch(r'[a-z0-9_-]{2,80}', x)))


def strict_title_match(title, roles):
    """Only direct target roles or a limited set of adjacent management titles."""
    if any(role_similarity(title, role) >= 3 for role in roles):
        return True
    value = title.lower()
    adjacent = [r'\b(?:director|head|vice president|vp) of (?:operations|business development|marketing|distribution|partnerships)\b',
                r'\b(?:regional|district|area) (?:manager|director)\b',
                r'\b(?:general manager|chief operating officer|coo)\b',
                r'\b(?:operations|business development|community engagement) director\b']
    return any(re.search(pattern, value) for pattern in adjacent)


def html_to_text(value):
    return re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', value))).strip()


def us_eligible(job):
    return bool(job.get('eligible', False))


ROLE_ALIASES = {
    'Chief Operating Officer (smaller organization)': ['chief operating officer', 'coo', 'head of operations', 'vp operations', 'vice president operations'],
    'Distribution Operations Director': ['distribution director', 'director of distribution', 'logistics director', 'director of logistics', 'supply chain director'],
    'Regional Operations Director': ['regional operations', 'regional director', 'director of operations', 'operations director', 'area manager'],
    'Digital Business Director': ['digital director', 'director of digital', 'digital business', 'digital strategy director', 'digital operations'],
    'Director of Community Partnerships': ['partnerships director', 'director of partnerships', 'community relations director', 'community engagement director'],
    'Director of Business Development': ['business development director', 'director of business development', 'head of business development', 'commercial director'],
    'Director of Marketing': ['marketing director', 'director of marketing', 'head of marketing'],
    'General Manager': ['general manager', 'business unit manager', 'managing director'],
}


def role_similarity(title, role):
    value = re.sub(r'[^a-z0-9 ]', ' ', title.lower())
    value = re.sub(r'\s+', ' ', value).strip()
    aliases = ROLE_ALIASES.get(role, [role.lower()])
    for phrase in aliases:
        clean = re.sub(r'[^a-z0-9 ]', ' ', phrase.lower())
        clean = re.sub(r'\s+', ' ', clean).strip()
        if re.search(r'\b' + re.escape(clean) + r'\b', value):
            return 3
    # Broader but still meaningful leadership-title signals.
    if role in ['Regional Operations Director', 'Distribution Operations Director'] and 'operations' in value and any(t in value for t in ['director', 'head', 'vp']):
        return 1
    return 0


# GovernmentJobs/NEOGOV supports agency-specific RSS feeds, not a single
# nationwide feed. Agency identifiers below were established from actual
# GovernmentJobs career-page URLs; availability of their RSS is checked live.
VERIFIED_NEOGOV_AGENCIES = {
    'hennepin': ('Hennepin County', 'MN'),
    'minneapolismn': ('City of Minneapolis', 'MN'),
    'kingcounty': ('King County', 'WA'),
}

# Links below lead to state employer career sites where known. The remaining
# states are covered by a 50-state public-sector search and an official
# government directory; no unverified state-agency URLs are invented.
STATE_CAREER_PORTALS = {
    'AL': 'https://www.personnel.alabama.gov/',
    'AZ': 'https://www.azstatejobs.gov/',
    'AR': 'https://arcareers.arkansas.gov/',
    'CA': 'https://calcareers.ca.gov/',
    'CO': 'https://www.colorado.gov/jobs',
    'FL': 'https://jobs.myflorida.com/',
    'GA': 'https://careers.georgia.gov/',
    'HI': 'https://jobs.hawaii.gov/',
    'MN': 'https://mn.gov/mmb/careers/search-for-jobs/',
    'NY': 'https://statejobs.ny.gov/',
    'NC': 'https://www.governmentjobs.com/careers/northcarolina',
    'OH': 'https://careers.ohio.gov/',
    'OR': 'https://oregon.wd5.myworkdayjobs.com/SOR_External_Career_Site',
    'PA': 'https://www.employment.pa.gov/',
    'TX': 'https://www.workintexas.com/',
    'VA': 'https://www.jobs.virginia.gov/',
    'WA': 'https://www.careers.wa.gov/',
    'WI': 'https://wisc.jobs/',
}


def parse_neogov_agencies(raw):
    """Accept agency tokens or GovernmentJobs agency career URLs only."""
    found = {}
    for token in re.split(r'[,\n;]+', raw or ''):
        token = token.strip()
        if not token:
            continue
        if token.startswith('http'):
            parsed = urlparse(token)
            if parsed.hostname not in ('governmentjobs.com', 'www.governmentjobs.com'):
                continue
            match = re.match(r'^/careers/([A-Za-z0-9_-]+)', parsed.path)
            if not match:
                continue
            agency = match.group(1).lower()
        else:
            agency = token.lower()
        if re.fullmatch(r'[a-z0-9_-]{2,70}', agency):
            found[agency] = VERIFIED_NEOGOV_AGENCIES.get(agency, (agency, ''))
    return dict(list(found.items())[:35])


@st.cache_data(ttl=3600, show_spinner=False)
def load_neogov_feeds(agencies):
    """Use documented NEOGOV RSS endpoint, never scrape career pages."""
    jobs, errors, counts = [], [], {}
    for agency, employer, state in agencies:
        url = 'https://www.governmentjobs.com/SearchEngine/JobsFeed?' + urlencode({'agency': agency})
        try:
            req = Request(url, headers={'User-Agent': 'AIJobFinder/1.0 (personal job search)', 'Accept': 'application/rss+xml, application/xml, text/xml'})
            with urlopen(req, timeout=10) as response:
                payload = response.read(2_000_000)
            root = ET.fromstring(payload)
            items = root.findall('.//item')
            if not items:
                # Some agencies don't expose RSS or have no current openings.
                errors.append(f'NEOGOV {employer}: feed returned no RSS items (check the agency career page).')
            count = 0
            for item in items[:250]:
                def field(name):
                    value = item.findtext(name)
                    return (value or '').strip()
                title = field('title')
                link = field('link')
                posted = field('pubDate') or field('date')
                description = field('description')
                # RSS does not consistently contain city/state. Use a verified
                # agency's known state only; do not assign states to user tokens.
                location = f'{employer}, {state}' if state else employer
                job = standard_job(title, employer, link, location, description,
                                   'NEOGOV RSS', remote=False, posted=posted)
                if job:
                    jobs.append(job)
                    count += 1
            counts[f'NEOGOV: {employer}'] = count
        except (HTTPError, URLError, TimeoutError, ET.ParseError, ValueError) as exc:
            errors.append(f'NEOGOV {employer}: feed unavailable ({type(exc).__name__}). Use its official careers page.')
            counts[f'NEOGOV: {employer}'] = 0
    return jobs, errors, counts


def job_competency_hits(job, resume):
    # Uses verified experience or phrases in the resume; compares to job description.
    confirmed = verified_competencies()
    if resume.strip():
        confirmed |= {item['skill'] for item in analyze_locally(resume)['documented']}
    job_text = job['title'] + ' ' + job['description']
    matches = []
    for area, phrases, *_ in COMPETENCIES:
        if area in confirmed and any(re.search(r'(?<!\w)' + re.escape(term) + r'(?!\w)', job_text, re.I) for term in phrases):
            matches.append(area)
    return matches[:5]


st.sidebar.title('💼 AI Job Finder')
page = st.sidebar.radio('Navigation', ['Dashboard','My Resume','Career Analysis','Career Matches','Find Jobs','Job Tracker'])
st.sidebar.caption('Free local resume comparison. No API key or AI usage charges.')

if page == 'Dashboard':
    st.title('📊 Job Search Dashboard')
    jobs = st.session_state.jobs
    a, b, c = st.columns(3)
    a.metric('Tracked Jobs', len(jobs))
    b.metric('Applied', sum(j['Status'] == 'Applied' for j in jobs))
    c.metric('Interviewing', sum(j['Status'] == 'Interviewing' for j in jobs))
    a, b, c = st.columns(3)
    a.metric('Offers', sum(j['Status'] == 'Offer' for j in jobs))
    b.metric('Saved', sum(j['Status'] == 'Saved' for j in jobs))
    c.metric('Skipped', sum(j['Status'] == 'Skipped' for j in jobs))
    if jobs:
        st.dataframe(pd.DataFrame(jobs), use_container_width=True, hide_index=True)
    else:
        st.info('Go to Job Tracker to add a job.')
    st.caption('Jobs are stored only for this session. Export your CSV to keep a backup.')

elif page == 'My Resume':
    st.title('📄 My Resume')
    file = st.file_uploader('Upload PDF, Word, or TXT resume', type=['pdf','docx','txt'])
    if file:
        try:
            extracted = extract_resume_text(file)
            if extracted.strip():
                if extracted != st.session_state.resume_text:
                    st.session_state.resume_text = extracted
                    st.session_state.analysis = None
                    st.session_state.pop('career_matches', None)
                    st.session_state.approved_suggestions = []
                st.success('Resume uploaded successfully!')
            else:
                st.warning('No readable text was found. Scanned PDFs may need OCR.')
        except Exception as exc:
            st.error(f'Unable to read file: {exc}')
    text = st.session_state.resume_text
    if text:
        with st.expander('View extracted resume text'):
            st.text_area('Resume text', text, height=280)
        st.subheader('Skills mentioned in your resume')
        matches = keyword_matches(text, SKILLS)
        st.write(', '.join(matches) if matches else 'No matches from the starter list.')
        st.caption('These are keyword matches, not verified skill assessments.')
        st.subheader('Confirm your actual job titles')
        found = keyword_matches(text, TITLES)
        if not st.session_state.confirmed_titles:
            st.session_state.confirmed_titles = ['Group Newspaper Publisher']
        choices = list(dict.fromkeys(found + st.session_state.confirmed_titles))
        selected = st.multiselect('Recognized or previously confirmed titles', choices, default=[t for t in st.session_state.confirmed_titles if t in choices])
        additional = st.text_input('Add other actual titles (comma-separated)')
        confirmed = list(dict.fromkeys(selected + [x.strip() for x in additional.split(',') if x.strip()]))
        st.session_state.confirmed_titles = confirmed
        st.write('**Confirmed titles:**', ', '.join(confirmed) if confirmed else 'None yet')
        st.info('Next, open Career Analysis to compare your experience against publisher responsibilities.')
    else:
        st.info('Upload a resume to begin.')

elif page == 'Career Analysis':
    st.title('🧭 Publisher Competency & Transferable Skills')
    st.write('Free, local comparison of your resume with common Group Newspaper Publisher responsibilities. Nothing is sent to an AI service.')
    actual_title = st.text_input('Your actual position', value='Group Newspaper Publisher')
    st.caption('Your real job title stays separate from transferable job categories.')
    verification_editor()
    if not st.session_state.resume_text:
        st.warning('First upload your resume on the My Resume page.')
    else:
        if st.button('Analyze My Resume', type='primary'):
            st.session_state.analysis = analyze_locally(st.session_state.resume_text)
            st.session_state.approved_suggestions = []
            st.success('Comparison complete. Review each result for accuracy.')
        analysis = st.session_state.analysis
        if analysis:
            st.subheader('Possible documented strengths')
            st.caption('A matching phrase is a clue, not proof you performed the duty. Review the source text.')
            for item in analysis['documented']:
                with st.expander(item['skill'], expanded=False):
                    st.write('**Resume text that matched:**')
                    for line in item['evidence']:
                        st.write(line)
                    st.write('**Transferable area:**', item['transferable'])
            st.subheader('Confirmed experience not necessarily documented in the resume')
            documented_skills = {item['skill'] for item in analysis['documented']}
            for duty in st.session_state.verified_duties:
                area = VERIFICATION_AREAS[duty]
                if area not in documented_skills:
                    st.write('✓ ' + duty + ' — **confirmed by you; add supporting details to your resume**')
            st.subheader('Other publisher competencies to investigate')
            missing = [area for area in analysis['verify'] if area not in verified_competencies()]
            if missing:
                for area in missing:
                    st.write('• ' + area)
            else:
                st.write('No additional competencies missing from this starter library.')
            st.subheader('Potential related career paths')
            for role, skills in analysis['roles'].items():
                st.write(f'**{role}** — related experience: {", ".join(skills)}')
            st.caption('These are ideas for exploration, not eligibility or job-match guarantees.')
            st.subheader('Draft wording for confirmed experience')
            st.caption('These are responsibility statements, not achievement claims. Add real scale, outcomes, and numbers only when accurate.')
            verified_drafts = []
            for i, duty in enumerate(st.session_state.verified_duties):
                proposed = st.text_area('Draft for: ' + duty,
                    value=duty.rstrip('.') + '.', key='verified_draft_' + str(i))
                if st.checkbox('Approve verified responsibility ' + str(i + 1), key='verified_approve_' + str(i)) and proposed.strip():
                    verified_drafts.append(proposed.strip())
            st.subheader('Resume wording suggestions — your approval required')
            st.write('These are editable starting points, not finished resume bullets.')
            approved = []
            for i, item in enumerate(analysis['suggestions']):
                st.markdown(f'**{i + 1}. {item["skill"]}**')
                st.caption('Original evidence: ' + item['source'])
                wording = st.text_area('Edit suggested wording', value=item['wording'], key=f'suggestion_{i}')
                if st.checkbox('Approve this wording', key=f'approve_{i}') and wording.strip():
                    approved.append(wording.strip())
            approved = list(dict.fromkeys(verified_drafts + approved))
            st.session_state.approved_suggestions = approved
            if approved:
                st.download_button('⬇️ Download approved wording', '\n\n'.join(approved), 'approved_resume_suggestions.txt', 'text/plain')
            st.info('Nothing is added to your resume automatically. You approve and download only the wording you want.')

elif page == 'Career Matches':
    st.title('🎯 Transferable Career Matches')
    st.write('Explore career paths outside newspaper publishing based on the responsibilities actually mentioned in your resume.')
    st.info('Free local matching combines resume phrase evidence with responsibilities you explicitly confirmed. Scores are not hiring predictions.')
    if not st.session_state.resume_text.strip():
        st.warning('First upload your resume on the My Resume page.')
    else:
        st.write('**Actual career background:** ' + (', '.join(st.session_state.confirmed_titles) or 'Group Newspaper Publisher'))
        st.caption('Your real title is not changed to match another industry.')
        st.write('**Confirmed responsibilities used in matching:**', len(st.session_state.verified_duties))
        st.caption('Edit these on Career Analysis → Experience Verification.')
        if st.button('Compare My Career Skills', type='primary'):
            st.session_state.career_matches = get_career_matches(st.session_state.resume_text)
        if 'career_matches' in st.session_state:
            results = st.session_state.career_matches
            industry_options = ['All industries'] + sorted({r['industry'] for r in results})
            industry = st.selectbox('Filter by industry', industry_options)
            visible = [r for r in results if industry == 'All industries' or r['industry'] == industry]
            st.caption('Overlap scores combine resume phrases and your confirmed responsibilities, weighted by career path. They are not hiring probabilities.')
            for r in visible:
                with st.expander(f"{r['title']} — {r['score']}% documented competency overlap", expanded=False):
                    st.write('**Industry:**', r['industry'])
                    st.progress(r['score'] / 100)
                    st.write('**Matched competency areas:**')
                    if r['matched']:
                        for skill in r['matched']:
                            st.markdown('**' + skill + '**')
                            if skill in r['evidence']:
                                st.caption('Resume phrase: ' + ' | '.join(r['evidence'][skill][:2]))
                            if skill in r['verified']:
                                st.caption('Confirmed by you: ' + '; '.join(d for d in st.session_state.verified_duties if VERIFICATION_AREAS[d] == skill))
                    else:
                        st.write('None from this starter competency list.')
                    st.write('**Responsibilities to confirm or document:**')
                    if r['unconfirmed']:
                        for skill in r['unconfirmed']:
                            st.write('• ' + skill)
                    else:
                        st.write('All mapped areas had phrase matches; verify that the duties were actually performed.')
                    st.write('**Resume tailoring idea:**', r['focus'])
                    st.caption('Only add achievements, responsibilities and numbers you can substantiate. Use Career Analysis to edit and approve suggested wording.')
            st.download_button('⬇️ Export Career Matches CSV',
                pd.DataFrame([{'Career':r['title'], 'Industry':r['industry'],
                    'Documented overlap (%)':r['score'],
                    'Matched areas':'; '.join(r['matched']),
                    'User-confirmed areas':'; '.join(r['verified']),
                    'Areas to verify':'; '.join(r['unconfirmed']),
                    'Resume focus':r['focus']} for r in visible]).to_csv(index=False),
                'career_matches.csv', 'text/csv')

elif page == 'Find Jobs':
    st.title('🔎 Find Real Jobs — Focused Search')
    st.write('Search U.S. location-tagged and remote job postings. State filters use the location provided by the employer; results are not comprehensive.')
    st.caption('Automatic sources: public job feeds, optional employer boards, USAJOBS, and selected agency-specific NEOGOV RSS feeds. Listings must be dated within 45 days.')
    recommended = [r['title'] for r in get_career_matches(st.session_state.resume_text)] if st.session_state.resume_text.strip() else [r['title'] for r in CAREER_PATHS]
    with st.form('live_job_search'):
        role = st.selectbox('Career path', ['All recommended career paths'] + recommended)
        keywords = st.text_input('Optional keywords (try operations, leadership, business development)')
        arrangement = st.selectbox('Work arrangement', ['On-site / hybrid only (hide remote)', 'Nationwide + remote', 'Remote only'], index=0)
        state_choice = st.selectbox('State', ['All states'] + [f'{name} ({code})' for code, name in sorted(US_STATES.items(), key=lambda item: item[1])])
        sort_choice = st.selectbox('Sort results', ['Best career match', 'State (A–Z)', 'Newest posting', 'Company (A–Z)'])
        strict_match = st.checkbox('Require a closely related management job title', value=True)
        exclude_companies = st.text_input('Exclude employers (comma-separated)', value='SpaceX', help='Example: SpaceX, Tesla')
        company_limit = st.slider('Maximum listings per employer', 1, 10, 3)
        max_results = st.slider('Maximum displayed jobs', 5, 60, 25, step=5)
        with st.expander('More sources — optional, no paid API required'):
            greenhouse_input = st.text_area('Greenhouse employer board tokens (comma-separated)', help='Use the employer token from boards.greenhouse.io/EMPLOYER. Not a general cross-employer search.')
            lever_input = st.text_area('Lever employer board tokens (comma-separated)', help='Use the employer token from jobs.lever.co/EMPLOYER.')
            st.caption('USAJOBS offers a free API key, but requires registration. Put USAJOBS_API_KEY and USAJOBS_EMAIL in Streamlit Secrets to enable automatic federal results.')
        with st.expander('Government RSS feeds — city and county employers'):
            use_gov_feeds = st.checkbox('Include NEOGOV government RSS feeds', value=True)
            gov_agencies = st.text_area('GovernmentJobs agency names or career URLs (comma/newline-separated)',
                value='hennepin, minneapolismn, kingcounty',
                help='Example: https://www.governmentjobs.com/careers/hennepin. Only entered agencies are checked; this is not a nationwide NEOGOV index.')
            st.caption('Agency feeds are checked individually, cached for one hour, and may be unavailable. Missing dates are excluded by the 45-day rule.')
        search = st.form_submit_button('Search Live Jobs', type='primary')
    if search:
        with st.spinner('Searching public job sources (may take up to a minute)...'):
            fetched, errors, counts = load_public_jobs()
            extra, extra_errors, extra_counts = load_employer_boards(
                parse_board_tokens(greenhouse_input), parse_board_tokens(lever_input))
            fetched.extend(extra)
            errors.extend(extra_errors)
            counts.update(extra_counts)
            try:
                federal_key = st.secrets.get('USAJOBS_API_KEY', '')
                federal_email = st.secrets.get('USAJOBS_EMAIL', '')
            except (FileNotFoundError, KeyError):
                federal_key, federal_email = '', ''
            federal, federal_errors, federal_counts = load_usajobs(federal_key, federal_email)
            fetched.extend(federal)
            errors.extend(federal_errors)
            counts.update(federal_counts)
            if use_gov_feeds:
                agency_map = parse_neogov_agencies(gov_agencies)
                agency_rows = tuple((token, name, state) for token, (name, state) in agency_map.items())
                govt, govt_errors, govt_counts = load_neogov_feeds(agency_rows)
                fetched.extend(govt)
                errors.extend(govt_errors)
                counts.update(govt_counts)
            fetched = list({j['url'].split('?')[0]: j for j in fetched}.values())
        st.session_state.feed_diagnostics = (len(fetched), counts, errors)
        selected_roles = recommended if role == 'All recommended career paths' else [role]
        state_code = state_choice.rsplit('(', 1)[-1].rstrip(')') if state_choice != 'All states' else None
        query = keywords.strip().lower()
        blocked = [name.strip().casefold() for name in exclude_companies.split(',') if name.strip()]
        filtered = []
        stale_count = 0
        unknown_date_count = 0
        for job in fetched:
            if not parse_posted_date(job.get('date')):
                unknown_date_count += 1
                continue
            if not is_recent_posting(job):
                stale_count += 1
                continue
            if not us_eligible(job):
                continue
            if any(name in job['company'].casefold() for name in blocked):
                continue
            if arrangement == 'On-site / hybrid only (hide remote)' and job['remote']:
                continue
            if arrangement == 'Remote only' and not job['remote']:
                continue
            job_states = states_in_location(job['location'])
            if state_code and state_code not in job_states:
                continue
            title = job['title'].lower()
            description = job['description'].lower()
            title_score = max((role_similarity(job['title'], r) for r in selected_roles), default=0)
            # Use transferable senior-management terms when the exact title differs.
            broad = bool(re.search(r'\b(operations|general manager|business development|partnerships|commercial|revenue|distribution|logistics|marketing|director|chief operating officer|publisher|regional manager)\b', title))
            extra = bool(query and query in (title + ' ' + description))
            if strict_match and not (strict_title_match(job['title'], selected_roles) or (extra and title_score)):
                continue
            if not (title_score or broad or extra):
                continue
            competency_hits = job_competency_hits(job, st.session_state.resume_text)
            rank = title_score * 4 + len(competency_hits) + (2 if extra else 0)
            filtered.append({**job, 'states': job_states, 'rank': rank, 'role_score': title_score, 'competency_hits': competency_hits})
        if sort_choice == 'State (A–Z)':
            filtered.sort(key=lambda j: (US_STATES.get(j['states'][0], 'ZZZ') if j['states'] else 'ZZZ', -j['rank']))
        elif sort_choice == 'Newest posting':
            filtered.sort(key=lambda j: (j['date'] or '0000-00-00', j['rank']), reverse=True)
        elif sort_choice == 'Company (A–Z)':
            filtered.sort(key=lambda j: (j['company'].lower(), -j['rank']))
        else:
            filtered.sort(key=lambda j: (-j['rank'], j['title'].lower()))
        # Diversity control runs before the final display cap.
        employer_counts = Counter()
        diverse = []
        for job in filtered:
            company_key = job['company'].casefold().strip()
            if employer_counts[company_key] < company_limit:
                diverse.append(job)
                employer_counts[company_key] += 1
        filtered = diverse
        st.session_state.date_filter_diagnostics = (stale_count, unknown_date_count)
        st.session_state.search_state_counts = dict(sorted(Counter(
            code for j in filtered for code in j['states']).items()))
        st.session_state.search_results = filtered[:max_results]
        st.session_state.search_total = len(filtered)
        st.session_state.search_timestamp = str(date.today())
    if 'feed_diagnostics' in st.session_state:
        count, counts, errors = st.session_state.feed_diagnostics
        with st.expander('Search diagnostics — see what sources returned'):
            st.write(f'{count} unique postings fetched before career filtering.')
            st.write('Source counts:', counts)
            stale, undated = st.session_state.get('date_filter_diagnostics', (0, 0))
            st.write(f'Excluded {stale} postings older than {MAX_POSTING_AGE_DAYS} days or dated in the future; excluded {undated} postings without a usable publication date.')
            for error in errors:
                st.warning(error)
            st.caption('If a source is blocked or unavailable, its count may be missing. Results are cached for one hour.')
    if 'search_results' in st.session_state:
        results = st.session_state.search_results
        st.subheader(f'{len(results)} jobs displayed')
        st.caption(f'Posting date filter: last {MAX_POSTING_AGE_DAYS} days only. Undated listings are hidden.')
        state_counts = st.session_state.get('search_state_counts', {})
        if state_counts:
            with st.expander('Job counts by state (before display limit)'):
                st.write(', '.join(f'{US_STATES.get(code, code)}: {count}' for code, count in state_counts.items()))
        else:
            st.caption('No explicit U.S. state locations in the current filtered results.')
        st.caption(f'{st.session_state.get("search_total", len(results))} filtered career-related postings • Searched {st.session_state.get("search_timestamp", "today")}. This is not a hiring-fit score.')
        if not results:
            st.info('No matching jobs for these filters in the available feeds. Try All states, Nationwide + remote, or a broader career path. Check Search diagnostics for source errors.')
        for i, job in enumerate(results):
            with st.container(border=True):
                st.markdown(f"**{job['title']}** — {job['company']}")
                st.write(f"📍 {job['location']} | {'Remote' if job['remote'] else 'On-site / hybrid (verify)' } | {job['source']}")
                if job['salary']:
                    st.write('**Salary (as listed):**', job['salary'])
                st.caption('Potential experience overlap: ' + (', '.join(job['competency_hits']) if job['competency_hits'] else 'No explicit competency phrase matches'))
                st.caption('Check the full posting for U.S. eligibility, work authorization, location restrictions and qualifications.')
                if job['date']:
                    st.caption('Posted: ' + job['date'])
                st.link_button('View original posting / Apply', job['url'])
                if st.button('Save to Job Tracker', key=f'save_live_{i}'):
                    if any(j.get('Job URL') == job['url'] for j in st.session_state.jobs):
                        st.info('Already in your Job Tracker.')
                    else:
                        st.session_state.jobs.append({
                            'Company': job['company'], 'Job Title': job['title'],
                            'Location': job['location'], 'Job URL': job['url'],
                            'Salary': job['salary'], 'Status': 'Saved',
                            'Date Recorded': str(date.today()), 'Interview Date': '',
                            'Notes': f"Found via {job['source']}. Confirm eligibility and requirements before applying."
                        })
                        st.success('Saved to Job Tracker for this session.')
    st.subheader('🏛️ Government jobs — all 50 states')
    st.caption('The state selector covers every U.S. state. Direct state-employer links are provided where configured; public-sector search links are available for all 50 states. External search results are NOT included in the 45-day filtered count; check dates and eligibility on the original website.')
    state_portal = st.selectbox('Select a state for government vacancies',
        [f'{name} ({code})' for code, name in sorted(US_STATES.items(), key=lambda kv: kv[1]) if code != 'DC'],
        index=next((i for i, (code, name) in enumerate(sorted(((c, n) for c, n in US_STATES.items() if c != 'DC'), key=lambda kv: kv[1])) if code == 'MN'), 0),
        key='government_state')
    selected_code = state_portal.rsplit('(', 1)[-1].rstrip(')')
    selected_name = US_STATES[selected_code]
    st.link_button(f'Search public-sector jobs in {selected_name}',
        'https://www.governmentjobs.com/jobs?' + urlencode({'location': selected_name}))
    if selected_code in STATE_CAREER_PORTALS:
        st.link_button(f'Open {selected_name} state government career portal', STATE_CAREER_PORTALS[selected_code])
    else:
        st.info('A direct state-employer portal has not been verified for this state in the app. Use the official 50-state directory to locate it; the public-sector search link above remains available.')
    st.link_button('Official 50-state government directory (USAGov)', 'https://www.usa.gov/state-governments')
    st.link_button('50-state career portal directory (UW–Madison)', 'https://career.lafollette.wisc.edu/resources/state-job-homepages-for-all-50-states/')
    st.link_button('Search federal government vacancies (USAJOBS)', 'https://www.usajobs.gov/Search/Results')
    st.link_button('Browse all local government jobs', 'https://www.governmentjobs.com/jobs')
    st.caption('NEOGOV auto-search works only for agency tokens entered in the search form, when those RSS feeds are accessible. It does not automatically cover every city or county nationwide.')
    st.caption('For automatic federal listings, add USAJOBS_API_KEY and USAJOBS_EMAIL to Streamlit Secrets after requesting a free USAJOBS API key. No OpenAI key is needed.')
    st.warning('Job Tracker entries are not permanently stored yet. Export your CSV before ending your session.')

elif page == 'Job Tracker':
    st.title('📋 Job Application Tracker')
    with st.expander('➕ Add a New Job', expanded=True):
        with st.form('add_job'):
            company = st.text_input('Company')
            title = st.text_input('Job Title')
            location = st.text_input('Location')
            url = st.text_input('Job Posting URL')
            salary = st.text_input('Salary / Range')
            status = st.selectbox('Status', STATUSES)
            recorded = st.date_input('Date Recorded', value=date.today())
            interview = st.date_input('Interview Date (optional)', value=None)
            notes = st.text_area('Notes / Next Steps')
            if st.form_submit_button('Save Job'):
                if company.strip() and title.strip():
                    st.session_state.jobs.append({'Company':company.strip(),'Job Title':title.strip(),'Location':location.strip(),'Job URL':url.strip(),'Salary':salary.strip(),'Status':status,'Date Recorded':str(recorded),'Interview Date':str(interview) if interview else '', 'Notes':notes.strip()})
                    st.success('Job added.')
                else:
                    st.warning('Company and Job Title are required.')
    jobs = st.session_state.jobs
    if jobs:
        selected_status = st.selectbox('Filter by status', ['All'] + STATUSES)
        visible = [j for j in jobs if selected_status == 'All' or j['Status'] == selected_status]
        st.dataframe(pd.DataFrame(visible), use_container_width=True, hide_index=True)
        idx = st.selectbox('Update a job', range(len(jobs)), format_func=lambda i: f"{jobs[i]['Company']} — {jobs[i]['Job Title']}")
        job = jobs[idx]
        new_status = st.selectbox('New status', STATUSES, index=STATUSES.index(job['Status']))
        new_notes = st.text_area('Update notes', value=job['Notes'], key=f'notes_{idx}')
        if st.button('Save Changes'):
            jobs[idx]['Status'] = new_status
            jobs[idx]['Notes'] = new_notes
            st.success('Changes saved for this session.')
            st.rerun()
        st.download_button('⬇️ Export Job Tracker CSV', pd.DataFrame(jobs).to_csv(index=False), 'job_tracker.csv', 'text/csv')
    else:
        st.info('No jobs tracked yet.')
    st.warning('Tracker data is temporary. Export a CSV backup before closing your session. Permanent storage is a future step.')
